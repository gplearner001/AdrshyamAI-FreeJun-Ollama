#!/usr/bin/env python3
"""
RAG Service for Knowledge Base Management
Handles document processing, chunking, embedding, and semantic search.
Uses PostgreSQL with pgvector for vector storage.
"""

import os
import logging
import io
import json
from typing import List, Dict, Any, Optional, BinaryIO
from datetime import datetime
import tiktoken
import psycopg2
from psycopg2.extras import RealDictCursor, Json
from psycopg2.pool import SimpleConnectionPool

try:
    import voyageai
    VOYAGE_AVAILABLE = True
except ImportError:
    VOYAGE_AVAILABLE = False
    voyageai = None

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format='%(levelname)s:%(name)s:%(message)s'
)
logger = logging.getLogger(__name__)

print("[RAG Service] Loading rag_service.py module...")
print(f"[RAG Service] VOYAGE_AVAILABLE: {VOYAGE_AVAILABLE}")

class RAGService:
    """Service for RAG-based knowledge base operations using PostgreSQL + pgvector."""

    def __init__(self):
        print("[RAG Service] Initializing RAGService...")
        self.voyage_api_key = os.getenv('VOYAGE_API_KEY')
        self.database_url = os.getenv('DATABASE_URL')

        self.voyage_client = None
        self.db_pool = None

        print(f"[RAG Service] Voyage API Key found: {bool(self.voyage_api_key)}")
        print(f"[RAG Service] Database URL found: {bool(self.database_url)}")

        logger.info(f"Voyage API Key found: {bool(self.voyage_api_key)}")
        logger.info(f"Database URL found: {bool(self.database_url)}")
        logger.info(f"Voyage library available: {VOYAGE_AVAILABLE}")

        if not self.voyage_api_key:
            logger.warning("VOYAGE_API_KEY environment variable not set!")
        if not self.database_url:
            logger.warning("DATABASE_URL environment variable not set!")

        if VOYAGE_AVAILABLE and self.voyage_api_key:
            try:
                print("[RAG Service] Attempting to initialize Voyage AI client...")
                self.voyage_client = voyageai.Client(api_key=self.voyage_api_key)
                print("[RAG Service] ✓ Voyage AI client initialized successfully")
                logger.info("Voyage AI client initialized successfully")
            except Exception as e:
                print(f"[RAG Service] ✗ Failed to initialize Voyage AI client: {str(e)}")
                logger.error(f"Failed to initialize Voyage AI client: {str(e)}")
        else:
            print(f"[RAG Service] ✗ Cannot initialize Voyage - Available: {VOYAGE_AVAILABLE}, Key: {bool(self.voyage_api_key)}")

        if self.database_url:
            try:
                print("[RAG Service] Attempting to initialize PostgreSQL connection pool...")
                self.db_pool = SimpleConnectionPool(1, 20, self.database_url)
                print("[RAG Service] ✓ PostgreSQL connection pool initialized successfully")
                logger.info("PostgreSQL connection pool initialized successfully")

                # Initialize database schema
                self._init_database_schema()
            except Exception as e:
                print(f"[RAG Service] ✗ Failed to initialize PostgreSQL connection: {str(e)}")
                logger.error(f"Failed to initialize PostgreSQL connection: {str(e)}")
        else:
            print(f"[RAG Service] ✗ Cannot initialize PostgreSQL - URL not provided")

        self.tokenizer = None
        try:
            self.tokenizer = tiktoken.encoding_for_model("gpt-4")
        except Exception as e:
            logger.warning(f"Failed to load tokenizer: {str(e)}")

        print(f"\n[RAG Service] ===== Initialization Complete =====")
        print(f"[RAG Service] Available: {self.is_available()}")
        print(f"[RAG Service]   - Voyage Client: {self.voyage_client is not None}")
        print(f"[RAG Service]   - DB Pool: {self.db_pool is not None}\n")

        logger.info(f"RAG Service initialized - Available: {self.is_available()}")
        logger.info(f"  - Voyage Client: {self.voyage_client is not None}")
        logger.info(f"  - DB Pool: {self.db_pool is not None}")

    def _init_database_schema(self):
        """Initialize database schema with pgvector extension and tables."""
        if not self.db_pool:
            return

        conn = None
        try:
            conn = self.db_pool.getconn()
            cur = conn.cursor()

            # Enable pgvector extension
            cur.execute("CREATE EXTENSION IF NOT EXISTS vector;")

            # Create knowledge_bases table
            cur.execute("""
                CREATE TABLE IF NOT EXISTS knowledge_bases (
                    id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                    name TEXT NOT NULL,
                    description TEXT,
                    user_id TEXT NOT NULL,
                    is_active BOOLEAN DEFAULT true,
                    created_at TIMESTAMPTZ DEFAULT NOW(),
                    updated_at TIMESTAMPTZ DEFAULT NOW()
                );
            """)

            # Create documents table
            cur.execute("""
                CREATE TABLE IF NOT EXISTS documents (
                    id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                    knowledge_base_id UUID NOT NULL REFERENCES knowledge_bases(id) ON DELETE CASCADE,
                    filename TEXT NOT NULL,
                    file_type TEXT NOT NULL,
                    file_size INTEGER,
                    processing_status TEXT DEFAULT 'pending',
                    error_message TEXT,
                    created_at TIMESTAMPTZ DEFAULT NOW(),
                    processed_at TIMESTAMPTZ
                );
            """)

            # Create document_chunks table with vector column
            cur.execute("""
                CREATE TABLE IF NOT EXISTS document_chunks (
                    id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                    document_id UUID NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    knowledge_base_id UUID NOT NULL REFERENCES knowledge_bases(id) ON DELETE CASCADE,
                    chunk_text TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL,
                    embedding VECTOR(1024),
                    metadata JSONB,
                    created_at TIMESTAMPTZ DEFAULT NOW()
                );
            """)

            # Create indexes for better performance
            cur.execute("""
                CREATE INDEX IF NOT EXISTS idx_documents_kb_id
                ON documents(knowledge_base_id);
            """)

            cur.execute("""
                CREATE INDEX IF NOT EXISTS idx_chunks_doc_id
                ON document_chunks(document_id);
            """)

            cur.execute("""
                CREATE INDEX IF NOT EXISTS idx_chunks_kb_id
                ON document_chunks(knowledge_base_id);
            """)

            # Create vector similarity search index
            cur.execute("""
                CREATE INDEX IF NOT EXISTS idx_chunks_embedding
                ON document_chunks USING ivfflat (embedding vector_cosine_ops)
                WITH (lists = 100);
            """)

            conn.commit()
            logger.info("Database schema initialized successfully")
        except Exception as e:
            if conn:
                conn.rollback()
            logger.error(f"Error initializing database schema: {str(e)}")
        finally:
            if conn:
                self.db_pool.putconn(conn)

    def get_connection(self):
        """Get a database connection from the pool."""
        if not self.db_pool:
            raise Exception("Database pool not initialized")
        return self.db_pool.getconn()

    def return_connection(self, conn):
        """Return a database connection to the pool."""
        if self.db_pool and conn:
            self.db_pool.putconn(conn)

    def is_available(self) -> bool:
        """Check if RAG service is available."""
        return self.voyage_client is not None and self.db_pool is not None

    def extract_text_from_file(self, file_content: BinaryIO, file_type: str) -> str:
        """Extract text content from various file formats."""
        try:
            file_type = file_type.lower()

            if file_type == 'txt':
                return file_content.read().decode('utf-8', errors='ignore')

            elif file_type == 'pdf' and PDF_AVAILABLE:
                pdf_reader = PyPDF2.PdfReader(file_content)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                return text

            elif file_type in ['doc', 'docx'] and DOCX_AVAILABLE:
                doc = DocxDocument(file_content)
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                return text

            elif file_type in ['xls', 'xlsx'] and XLSX_AVAILABLE:
                workbook = openpyxl.load_workbook(file_content, read_only=True)
                text = ""
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                        text += row_text + "\n"
                return text

            else:
                logger.warning(f"Unsupported file type: {file_type}")
                return ""

        except Exception as e:
            logger.error(f"Error extracting text from {file_type}: {str(e)}")
            return ""

    def chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 200) -> List[Dict[str, Any]]:
        """Split text into overlapping chunks."""
        try:
            if not text or not text.strip():
                return []

            if self.tokenizer:
                tokens = self.tokenizer.encode(text)
                chunks = []

                for i in range(0, len(tokens), chunk_size - overlap):
                    chunk_tokens = tokens[i:i + chunk_size]
                    chunk_text = self.tokenizer.decode(chunk_tokens)

                    chunks.append({
                        'text': chunk_text,
                        'index': len(chunks),
                        'token_count': len(chunk_tokens)
                    })

                return chunks
            else:
                words = text.split()
                chunks = []

                for i in range(0, len(words), chunk_size - overlap):
                    chunk_words = words[i:i + chunk_size]
                    chunk_text = " ".join(chunk_words)

                    chunks.append({
                        'text': chunk_text,
                        'index': len(chunks),
                        'word_count': len(chunk_words)
                    })

                return chunks

        except Exception as e:
            logger.error(f"Error chunking text: {str(e)}")
            return []

    async def generate_embedding(self, text: str, input_type: str = "document") -> Optional[List[float]]:
        """Generate embedding vector for text using Voyage AI (Anthropic's recommended partner).

        Args:
            text: The text to embed
            input_type: Either "document" for indexing or "query" for searching
        """
        if not self.voyage_client:
            return None

        try:
            result = self.voyage_client.embed(
                texts=[text],
                model="voyage-2",
                input_type=input_type
            )
            return result.embeddings[0]

        except Exception as e:
            logger.error(f"Error generating embedding: {str(e)}")
            return None

    async def process_document(
        self,
        document_id: str,
        file_content: BinaryIO,
        file_type: str,
        knowledge_base_id: str
    ) -> Dict[str, Any]:
        """Process a document: extract text, chunk, and generate embeddings."""
        conn = None
        try:
            text = self.extract_text_from_file(file_content, file_type)

            if not text or not text.strip():
                return {
                    'success': False,
                    'error': 'Failed to extract text from document',
                    'chunks_created': 0
                }

            chunks = self.chunk_text(text)

            if not chunks:
                return {
                    'success': False,
                    'error': 'Failed to chunk document',
                    'chunks_created': 0
                }

            chunks_created = 0
            conn = self.get_connection()
            cur = conn.cursor()

            for chunk in chunks:
                embedding = await self.generate_embedding(chunk['text'])

                if embedding:
                    cur.execute("""
                        INSERT INTO document_chunks
                        (document_id, knowledge_base_id, chunk_text, chunk_index, embedding, metadata)
                        VALUES (%s, %s, %s, %s, %s, %s)
                    """, (
                        document_id,
                        knowledge_base_id,
                        chunk['text'],
                        chunk['index'],
                        embedding,
                        Json({
                            'token_count': chunk.get('token_count', 0),
                            'word_count': chunk.get('word_count', 0)
                        })
                    ))
                    chunks_created += 1

            cur.execute("""
                UPDATE documents
                SET processing_status = 'completed', processed_at = NOW()
                WHERE id = %s
            """, (document_id,))

            conn.commit()

            return {
                'success': True,
                'chunks_created': chunks_created,
                'total_chunks': len(chunks)
            }

        except Exception as e:
            if conn:
                conn.rollback()
            logger.error(f"Error processing document: {str(e)}")

            try:
                conn2 = self.get_connection()
                cur2 = conn2.cursor()
                cur2.execute("""
                    UPDATE documents
                    SET processing_status = 'failed', error_message = %s
                    WHERE id = %s
                """, (str(e), document_id))
                conn2.commit()
                self.return_connection(conn2)
            except:
                pass

            return {
                'success': False,
                'error': str(e),
                'chunks_created': 0
            }
        finally:
            if conn:
                self.return_connection(conn)

    async def search_knowledge_base(
        self,
        query: str,
        knowledge_base_id: str,
        limit: int = 5,
        threshold: float = 0.5
    ) -> List[Dict[str, Any]]:
        """Search knowledge base using semantic similarity."""
        logger.info(f"🔍 Searching KB {knowledge_base_id} for query: '{query}' (limit: {limit}, threshold: {threshold})")

        if not self.is_available():
            logger.warning("⚠️ RAG service not available for search")
            return []

        conn = None
        try:
            logger.info(f"🎯 Generating query embedding...")
            query_embedding = await self.generate_embedding(query, input_type="query")

            if not query_embedding:
                logger.error("❌ Failed to generate query embedding")
                return []

            logger.info(f"✓ Generated embedding with {len(query_embedding)} dimensions")

            conn = self.get_connection()
            cur = conn.cursor(cursor_factory=RealDictCursor)

            # Use cosine similarity for vector search
            logger.info(f"🔍 Executing vector similarity search...")
            cur.execute("""
                SELECT
                    chunk_text,
                    chunk_index,
                    document_id,
                    metadata,
                    1 - (embedding <=> %s::vector) as similarity
                FROM document_chunks
                WHERE knowledge_base_id = %s
                    AND 1 - (embedding <=> %s::vector) > %s
                ORDER BY embedding <=> %s::vector
                LIMIT %s
            """, (query_embedding, knowledge_base_id, query_embedding, threshold, query_embedding, limit))

            results = cur.fetchall()
            logger.info(f"✓ Found {len(results)} results above similarity threshold {threshold}")

            if results:
                for idx, result in enumerate(results):
                    logger.debug(f"  Result {idx + 1}: similarity={result['similarity']:.3f}, chunk_index={result.get('chunk_index')}")
            else:
                logger.warning(f"⚠️ No results found above similarity threshold {threshold}")

            return [dict(row) for row in results]

        except Exception as e:
            logger.error(f"❌ Error searching knowledge base: {str(e)}")
            import traceback
            logger.error(traceback.format_exc())
            return []
        finally:
            if conn:
                self.return_connection(conn)

    async def get_context_for_query(
        self,
        query: str,
        knowledge_base_id: str,
        max_tokens: int = 2000
    ) -> str:
        """Get relevant context from knowledge base for a query."""
        logger.info(f"🔍 Getting context for query: '{query}' from KB: {knowledge_base_id}")
        search_results = await self.search_knowledge_base(query, knowledge_base_id)

        logger.info(f"📊 Search returned {len(search_results)} results")

        if not search_results:
            logger.warning(f"⚠️ No search results found for query in KB {knowledge_base_id}")
            return ""

        context_parts = []
        current_tokens = 0

        for idx, result in enumerate(search_results):
            chunk_text = result.get('chunk_text', '')
            similarity = result.get('similarity', 0)

            if self.tokenizer:
                chunk_tokens = len(self.tokenizer.encode(chunk_text))
            else:
                chunk_tokens = len(chunk_text.split())

            if current_tokens + chunk_tokens > max_tokens:
                logger.info(f"📊 Token limit reached at result {idx + 1}/{len(search_results)}")
                break

            context_parts.append(f"[Relevance: {similarity:.2f}]\n{chunk_text}")
            current_tokens += chunk_tokens
            logger.debug(f"✓ Added result {idx + 1} (similarity: {similarity:.2f}, tokens: {chunk_tokens})")

        final_context = "\n\n---\n\n".join(context_parts)
        logger.info(f"✓ Built context with {len(context_parts)} chunks, {current_tokens} tokens, {len(final_context)} chars")
        return final_context

rag_service = RAGService()
